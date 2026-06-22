from __future__ import annotations

import math
import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import openpyxl
import pandas as pd
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel, Field, ValidationError


BASE_DIR = Path(__file__).resolve().parent
MODEL_PATH = BASE_DIR / "models" / "production_model.joblib"
PRICE_MODELS_PATH = BASE_DIR / "models" / "production_price_models.joblib"
TRUCK_WORKBOOK_PATH = BASE_DIR / "data" / "data.xlsx"
LOOKUP_SHEET_NAME = "Sheet1"
COORDINATE_PRECISION = 2

INDIAN_STATES = [
    "andhra pradesh",
    "arunachal pradesh",
    "assam",
    "bihar",
    "chhattisgarh",
    "delhi",
    "goa",
    "gujarat",
    "haryana",
    "himachal pradesh",
    "jharkhand",
    "karnataka",
    "kerala",
    "madhya pradesh",
    "maharashtra",
    "manipur",
    "meghalaya",
    "mizoram",
    "nagaland",
    "odisha",
    "orissa",
    "punjab",
    "rajasthan",
    "sikkim",
    "tamil nadu",
    "telangana",
    "tripura",
    "uttar pradesh",
    "uttarakhand",
    "west bengal",
    "andaman and nicobar islands",
    "chandigarh",
    "dadra and nagar haveli",
    "daman and diu",
    "jammu and kashmir",
    "ladakh",
    "lakshadweep",
    "puducherry",
]
STATE_ALIASES = {"pondicherry": "puducherry", "orissa": "odisha"}


class PricingError(ValueError):
    pass


class PriceRequest(BaseModel):
    pickup_latitude: float = Field(..., ge=-90, le=90)
    pickup_longitude: float = Field(..., ge=-180, le=180)
    drop_latitude: float = Field(..., ge=-90, le=90)
    drop_longitude: float = Field(..., ge=-180, le=180)
    pickup_location: str | None = None
    drop_location: str | None = None
    pickup_city: str | None = None
    destination_city: str | None = None
    pickup_state: str | None = None
    drop_state: str | None = None
    truck_type: str = Field(..., min_length=1)
    body_type: str = Field(..., min_length=1)
    weight: str = Field(..., min_length=1)
    distance_km: float = Field(..., gt=0)
    diesel_price: float = Field(..., gt=0)


class BulkPriceRequest(BaseModel):
    items: list[Any] = Field(..., min_length=1)


app = FastAPI(
    title="Route Price API",
    description="Predict budget, standard, and premium route prices from saved production models.",
    version="1.0.0",
)


def normalize_text_value(value: object) -> object:
    if pd.isna(value):
        return value
    return re.sub(r"\s+", " ", str(value).strip())


def normalize_vehicle_type(value: object) -> object:
    text = normalize_text_value(value)
    if pd.isna(text):
        return text
    text = str(text).lower()
    text = re.sub(r"\bfeet\b", "Feet", text)
    text = re.sub(r"\bfoot\b", "Feet", text)
    text = re.sub(r"\bft\b", "Feet", text)
    text = re.sub(r"\bwheel\b", "Wheel", text)
    text = re.sub(r"\btrailor\b", "Trailor", text)
    text = re.sub(r"\btrailer\b", "Trailor", text)
    text = re.sub(r"\btata ace\b", "Tata ace", text)
    text = re.sub(r"\bbada dost\b", "Bada dost", text)
    text = re.sub(r"\bdost\b", "Dost", text)
    text = re.sub(r"\bbolero\b", "Bolero", text)
    text = re.sub(r"\bjcb\b", "JCB", text)
    text = re.sub(r"\bmxl\b", "MXL", text)
    text = re.sub(r"\bsxl\b", "SXL", text)
    text = re.sub(r"\bhigh bed\b", "High Bed", text)
    text = re.sub(r"\bsemi bed\b", "Semi Bed", text)
    text = re.sub(r"\s*-\s*", " - ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_body_type(value: object) -> object:
    text = normalize_text_value(value)
    return text if pd.isna(text) else str(text).lower().title()


def normalize_load_size(value: object) -> object:
    text = normalize_text_value(value)
    if pd.isna(text):
        return text
    text = str(text).lower()
    text = re.sub(r"\btonnes?\b", "tons", text)
    text = re.sub(r"\btons?\b", "tons", text)
    text = re.sub(r"\bkgs?\b", "kg", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_route_part(value: object) -> str:
    text = str(value).lower().strip()
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or "unknown"


def coordinate_route_part(latitude: object, longitude: object, precision: int = COORDINATE_PRECISION) -> str:
    lat = pd.to_numeric(pd.Series([latitude]), errors="coerce").iloc[0]
    lon = pd.to_numeric(pd.Series([longitude]), errors="coerce").iloc[0]
    if pd.isna(lat) or pd.isna(lon):
        return "unknown"
    return f"{float(lat):.{precision}f},{float(lon):.{precision}f}"


def extract_state(location: str) -> str:
    normalized = str(location).lower()
    for state in INDIAN_STATES:
        if re.search(rf"\b{re.escape(state)}\b", normalized):
            return STATE_ALIASES.get(state, state)
    return "unknown"


def normalize_categorical_columns(df: pd.DataFrame) -> pd.DataFrame:
    string_columns = df.select_dtypes(include=["object", "string"]).columns
    for column in string_columns:
        df[column] = df[column].map(normalize_text_value)
    if "vehicle_type" in df.columns:
        df["vehicle_type"] = df["vehicle_type"].map(normalize_vehicle_type)
    if "body_type" in df.columns:
        df["body_type"] = df["body_type"].map(normalize_body_type)
    if "load_size" in df.columns:
        df["load_size"] = df["load_size"].map(normalize_load_size)
    return df


def add_features(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if {"from_location", "to_location"}.issubset(df.columns):
        origin_location = df["from_location"].fillna("").astype(str).str.lower().str.strip()
        destination_location = df["to_location"].fillna("").astype(str).str.lower().str.strip()
        df["route_id"] = origin_location + "__" + destination_location
        df["route_origin_state"] = origin_location.map(extract_state)
        df["route_destination_state"] = destination_location.map(extract_state)
        df["route_id_state"] = df["route_origin_state"] + "__" + df["route_destination_state"]

    if {"pickup_city", "destination_city"}.issubset(df.columns):
        df["route_origin_city"] = df["pickup_city"].fillna("").map(normalize_route_part)
        df["route_destination_city"] = df["destination_city"].fillna("").map(normalize_route_part)
        df["route_id_city"] = df["route_origin_city"] + "__" + df["route_destination_city"]

    coordinate_columns = {
        "pickup_latitude",
        "pickup_longitude",
        "drop_latitude",
        "drop_longitude",
    }
    if coordinate_columns.issubset(df.columns):
        for column in coordinate_columns:
            df[column] = pd.to_numeric(df[column], errors="coerce")
        df["route_latitude_delta"] = df["drop_latitude"] - df["pickup_latitude"]
        df["route_longitude_delta"] = df["drop_longitude"] - df["pickup_longitude"]
        origin_coordinate = [
            coordinate_route_part(lat, lon)
            for lat, lon in zip(df["pickup_latitude"], df["pickup_longitude"])
        ]
        destination_coordinate = [
            coordinate_route_part(lat, lon)
            for lat, lon in zip(df["drop_latitude"], df["drop_longitude"])
        ]
        df["route_id_coordinate"] = [
            f"{origin}__{destination}"
            for origin, destination in zip(origin_coordinate, destination_coordinate)
        ]
    return df


def _require_positive_number(value: float, field_name: str) -> float:
    if value is None or not math.isfinite(float(value)) or float(value) <= 0:
        raise PricingError(f"{field_name} must be a positive number.")
    return float(value)


def _require_coordinate(value: float, field_name: str, minimum: float, maximum: float) -> float:
    if value is None or not math.isfinite(float(value)):
        raise PricingError(f"{field_name} must be a valid coordinate.")
    number = float(value)
    if number < minimum or number > maximum:
        raise PricingError(f"{field_name} must be between {minimum} and {maximum}.")
    return number


def _normalize_vehicle_type_required(value: str) -> str:
    normalized = normalize_vehicle_type(value)
    if pd.isna(normalized):
        raise PricingError("truck_type is required.")
    return str(normalized)


def _google_maps_api_key() -> str | None:
    return os.getenv("GOOGLE_MAPS_API_KEY") or None


def _component(address_components: list[dict[str, Any]], *types: str) -> str:
    wanted = set(types)
    for component in address_components:
        if wanted.intersection(component.get("types", [])):
            return str(component.get("long_name", ""))
    return ""


def _place_payload_to_location(result: dict[str, Any]) -> dict[str, Any]:
    geometry = result.get("geometry") or {}
    point = geometry.get("location") or {}
    components = result.get("address_components") or []
    city = (
        _component(components, "locality")
        or _component(components, "postal_town")
        or _component(components, "administrative_area_level_3")
        or _component(components, "administrative_area_level_2")
    )
    return {
        "address": result.get("formatted_address") or result.get("name") or "",
        "city": city,
        "state": _component(components, "administrative_area_level_1"),
        "latitude": float(point["lat"]) if "lat" in point else None,
        "longitude": float(point["lng"]) if "lng" in point else None,
    }


@lru_cache(maxsize=256)
def reverse_geocode(latitude: float, longitude: float) -> dict[str, Any]:
    api_key = _google_maps_api_key()
    if not api_key:
        raise RuntimeError("Set GOOGLE_MAPS_API_KEY before calling Google Maps APIs.")

    import requests

    response = requests.get(
        "https://maps.googleapis.com/maps/api/geocode/json",
        params={"latlng": f"{latitude},{longitude}", "key": api_key},
        timeout=20,
    )
    response.raise_for_status()
    payload = response.json()
    status = payload.get("status")
    if status not in {"OK", "ZERO_RESULTS"}:
        raise RuntimeError(payload.get("error_message") or status or "Unknown Google Maps API error")
    results = payload.get("results") or []
    if not results:
        return {"address": "", "city": "", "state": "", "latitude": latitude, "longitude": longitude}
    return _place_payload_to_location(results[0])


def metadata_from_coordinates(
    latitude: float,
    longitude: float,
    city: str | None,
    state: str | None,
    location: str | None,
) -> dict[str, str]:
    if city and state:
        return {
            "location": f"{city}, {state}",
            "city": city,
            "state": state,
            "display_location": location or f"{city}, {state}",
        }

    try:
        metadata = reverse_geocode(latitude, longitude)
    except Exception:
        metadata = {}

    resolved_city = city or metadata.get("city") or "unknown"
    resolved_state = state or metadata.get("state") or "unknown"
    return {
        "location": f"{resolved_city}, {resolved_state}",
        "city": resolved_city,
        "state": resolved_state,
        "display_location": location or metadata.get("address") or f"{latitude},{longitude}",
    }


@lru_cache(maxsize=1)
def load_truck_calculation_master() -> dict[str, dict[str, Any]]:
    if not TRUCK_WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Truck lookup workbook not found: {TRUCK_WORKBOOK_PATH}")

    workbook = openpyxl.load_workbook(TRUCK_WORKBOOK_PATH, data_only=True, read_only=True)
    if LOOKUP_SHEET_NAME not in workbook.sheetnames:
        raise PricingError(f"Missing sheet: {LOOKUP_SHEET_NAME}")

    sheet = workbook[LOOKUP_SHEET_NAME]
    master: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        truck_type = row[0]
        if not truck_type:
            continue

        normalized_type = _normalize_vehicle_type_required(str(truck_type))
        master[normalized_type] = {
            "truck_type": normalized_type,
            "emi_per_day": _require_positive_number(row[1], "EMI/DAY"),
            "kmpl": _require_positive_number(row[2], "KMPL"),
            "wear_and_tear_rate_per_km": _require_positive_number(row[3], "WEAR AND TEAR"),
            "km_per_day": _require_positive_number(row[4], "KM/DAY"),
            "external_toll_rate_per_km": _require_positive_number(row[5], "TOLL"),
            "toll_guru_class": row[6] if len(row) > 6 else None,
        }
    workbook.close()

    if not master:
        raise PricingError("No truck calculation master rows were loaded.")
    return master


@lru_cache(maxsize=1)
def load_production_model():
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Production model not found: {MODEL_PATH}")
    return joblib.load(MODEL_PATH)


@lru_cache(maxsize=1)
def load_production_price_models() -> dict[str, dict[str, Any]]:
    if PRICE_MODELS_PATH.exists():
        return joblib.load(PRICE_MODELS_PATH)

    fallback_model = load_production_model()
    return {
        "standard": {
            "label": "Standard",
            "description": "Expected-price option from the production route-cost model.",
            "pipeline": fallback_model,
        }
    }


def pipeline_feature_columns(pipeline) -> tuple[list[str], list[str]]:
    preprocessor = pipeline.named_steps["preprocessor"]
    numeric_features: list[str] = []
    categorical_features: list[str] = []
    for name, _, columns in preprocessor.transformers:
        values = list(columns)
        if name == "num":
            numeric_features = values
        elif name == "cat":
            categorical_features = values
    return numeric_features, categorical_features


def prepare_feature_frame(
    df: pd.DataFrame,
    numeric_features: list[str],
    categorical_features: list[str],
) -> pd.DataFrame:
    frame = df[numeric_features + categorical_features].copy()
    for column in numeric_features:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    for column in categorical_features:
        frame[column] = frame[column].astype("object").where(frame[column].notna(), np.nan)
    return frame


def model_input_frame(
    pickup_latitude: float,
    pickup_longitude: float,
    drop_latitude: float,
    drop_longitude: float,
    pickup_location: str,
    drop_location: str,
    pickup_city: str,
    destination_city: str,
    truck_type: str,
    body_type: str,
    weight: str,
    distance_km: float,
    transit_days: int,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "from_location": pickup_location,
                "to_location": drop_location,
                "pickup_latitude": pickup_latitude,
                "pickup_longitude": pickup_longitude,
                "drop_latitude": drop_latitude,
                "drop_longitude": drop_longitude,
                "pickup_city": pickup_city,
                "destination_city": destination_city,
                "vehicle_type": truck_type,
                "body_type": body_type,
                "load_size": weight,
                "distance_km": distance_km,
                "transit_days": transit_days,
            }
        ]
    )


def predict_route_cost_options(
    pickup_latitude: float,
    pickup_longitude: float,
    drop_latitude: float,
    drop_longitude: float,
    pickup_location: str,
    drop_location: str,
    pickup_city: str,
    destination_city: str,
    truck_type: str,
    body_type: str,
    weight: str,
    distance_km: float,
    transit_days: int,
) -> dict[str, dict[str, Any]]:
    raw = model_input_frame(
        pickup_latitude,
        pickup_longitude,
        drop_latitude,
        drop_longitude,
        pickup_location,
        drop_location,
        pickup_city,
        destination_city,
        truck_type,
        body_type,
        weight,
        distance_km,
        transit_days,
    )
    featured = add_features(normalize_categorical_columns(raw))

    options = {}
    for name, model_entry in load_production_price_models().items():
        pipeline = model_entry["pipeline"]
        numeric_features, categorical_features = pipeline_feature_columns(pipeline)
        prepared = prepare_feature_frame(featured, numeric_features, categorical_features)
        prediction = pipeline.predict(prepared)[0]
        options[name] = {
            "label": model_entry.get("label", name.title()),
            "description": model_entry.get("description", ""),
            "route_cost": float(prediction),
        }
    enforce_price_option_order(options)
    return options


def enforce_price_option_order(options: dict[str, dict[str, Any]]) -> None:
    required_options = {"budget", "standard", "premium"}
    if not required_options.issubset(options):
        return

    ordered_costs = sorted(
        [
            options["budget"]["route_cost"],
            options["standard"]["route_cost"],
            options["premium"]["route_cost"],
        ]
    )
    options["budget"]["route_cost"] = ordered_costs[0]
    options["standard"]["route_cost"] = ordered_costs[1]
    options["premium"]["route_cost"] = ordered_costs[2]


def calculate_price_option(
    option: dict[str, Any],
    fuel_cost: float,
    emi_cost: float,
    wear_and_tear_cost: float,
    calculated_toll_cost: float,
) -> dict[str, Any]:
    route_cost = option["route_cost"]
    total_rate_85_percent = (
        fuel_cost + emi_cost + wear_and_tear_cost + calculated_toll_cost + route_cost
    )
    driver_rate = total_rate_85_percent / 0.85
    driver_salary = driver_rate * 0.15
    return {
        "label": option["label"],
        "description": option["description"],
        "route_cost": round(route_cost, 2),
        "total_rate_85_percent": round(total_rate_85_percent, 2),
        "driver_salary_15_percent": round(driver_salary, 2),
        "driver_rate": round(driver_rate, 2),
    }


def calculate_driver_rate(
    pickup_latitude: float,
    pickup_longitude: float,
    drop_latitude: float,
    drop_longitude: float,
    truck_type: str,
    body_type: str,
    weight: str,
    distance_km: float,
    diesel_price: float,
    pickup_location: str | None = None,
    drop_location: str | None = None,
    pickup_city: str | None = None,
    destination_city: str | None = None,
    pickup_state: str | None = None,
    drop_state: str | None = None,
) -> dict[str, Any]:
    pickup_latitude = _require_coordinate(pickup_latitude, "pickup_latitude", -90, 90)
    pickup_longitude = _require_coordinate(pickup_longitude, "pickup_longitude", -180, 180)
    drop_latitude = _require_coordinate(drop_latitude, "drop_latitude", -90, 90)
    drop_longitude = _require_coordinate(drop_longitude, "drop_longitude", -180, 180)
    distance_km = _require_positive_number(distance_km, "distance_km")
    diesel_price = _require_positive_number(diesel_price, "diesel_price")
    normalized_truck_type = _normalize_vehicle_type_required(truck_type)
    normalized_body_type = normalize_body_type(body_type)
    normalized_weight = normalize_load_size(weight)

    master = load_truck_calculation_master()
    if normalized_truck_type not in master:
        available = ", ".join(sorted(master))
        raise PricingError(
            f"Unsupported truck_type '{truck_type}'. Available truck types: {available}"
        )

    truck_master = master[normalized_truck_type]
    transit_days = int(math.ceil(distance_km / truck_master["km_per_day"]))
    pickup_metadata = metadata_from_coordinates(
        pickup_latitude,
        pickup_longitude,
        pickup_city,
        pickup_state,
        pickup_location,
    )
    drop_metadata = metadata_from_coordinates(
        drop_latitude,
        drop_longitude,
        destination_city,
        drop_state,
        drop_location,
    )
    route_cost_options = predict_route_cost_options(
        pickup_latitude,
        pickup_longitude,
        drop_latitude,
        drop_longitude,
        pickup_metadata["location"],
        drop_metadata["location"],
        pickup_metadata["city"],
        drop_metadata["city"],
        normalized_truck_type,
        str(normalized_body_type),
        str(normalized_weight),
        distance_km,
        transit_days,
    )

    fuel_cost = diesel_price * distance_km / truck_master["kmpl"]
    emi_cost = transit_days * truck_master["emi_per_day"]
    wear_and_tear_cost = truck_master["wear_and_tear_rate_per_km"] * distance_km
    calculated_toll_cost = truck_master["external_toll_rate_per_km"] * distance_km
    price_options = {
        name: calculate_price_option(
            option,
            fuel_cost,
            emi_cost,
            wear_and_tear_cost,
            calculated_toll_cost,
        )
        for name, option in route_cost_options.items()
    }
    standard_option = price_options.get("standard") or next(iter(price_options.values()))

    return {
        "input": {
            "pickup_latitude": pickup_latitude,
            "pickup_longitude": pickup_longitude,
            "drop_latitude": drop_latitude,
            "drop_longitude": drop_longitude,
            "pickup_location": pickup_metadata["display_location"],
            "drop_location": drop_metadata["display_location"],
            "pickup_city": pickup_metadata["city"],
            "destination_city": drop_metadata["city"],
            "pickup_state": pickup_metadata["state"],
            "drop_state": drop_metadata["state"],
            "truck_type": normalized_truck_type,
            "body_type": str(normalized_body_type),
            "weight": str(normalized_weight),
            "distance_km": distance_km,
            "diesel_price": diesel_price,
        },
        "calculation_master": truck_master,
        "days": transit_days,
        "fuel_cost_inr": round(fuel_cost, 2),
        "emi_cost": round(emi_cost, 2),
        "total_wear_and_tear_cost_inr": round(wear_and_tear_cost, 2),
        "calculated_toll_cost_inr": round(calculated_toll_cost, 2),
        "route_cost": standard_option["route_cost"],
        "total_rate_85_percent": standard_option["total_rate_85_percent"],
        "driver_salary_15_percent": standard_option["driver_salary_15_percent"],
        "driver_rate": standard_option["driver_rate"],
        "price_options": price_options,
    }


def model_payload(model: BaseModel) -> dict:
    return model.model_dump() if hasattr(model, "model_dump") else model.dict()


def validation_errors(error: ValidationError) -> list[dict]:
    return [
        {
            "field": ".".join(str(part) for part in item.get("loc", [])),
            "message": item.get("msg", "Invalid value"),
            "type": item.get("type", "validation_error"),
        }
        for item in error.errors()
    ]


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/predict-price")
def predict_price(request: PriceRequest) -> dict:
    try:
        return calculate_driver_rate(**model_payload(request))
    except (FileNotFoundError, PricingError) as error:
        raise HTTPException(status_code=400, detail=str(error)) from error


@app.post("/predict-prices")
def predict_prices(request: BulkPriceRequest) -> dict:
    results = []
    success_count = 0
    error_count = 0

    for index, item in enumerate(request.items):
        try:
            if not isinstance(item, dict):
                raise PricingError("Each bulk item must be a JSON object.")
            result = calculate_driver_rate(**model_payload(PriceRequest(**item)))
            results.append({"index": index, "status": "success", "result": result})
            success_count += 1
        except ValidationError as error:
            results.append(
                {
                    "index": index,
                    "status": "error",
                    "error_type": "validation_error",
                    "errors": validation_errors(error),
                }
            )
            error_count += 1
        except (FileNotFoundError, PricingError) as error:
            results.append(
                {
                    "index": index,
                    "status": "error",
                    "error_type": "pricing_error",
                    "message": str(error),
                }
            )
            error_count += 1

    return {
        "total": len(request.items),
        "success_count": success_count,
        "error_count": error_count,
        "results": results,
    }
